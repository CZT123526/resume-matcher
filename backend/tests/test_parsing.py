"""S01 文档解析测试。"""
import pytest
from io import BytesIO

from app.core.errors import AppError, ErrorCode
from app.svc.parsing import parse_document


def test_unsupported_ext_rejected():
    with pytest.raises(AppError) as ei:
        parse_document(b"x", "resume.txt")
    assert ei.value.code == ErrorCode.UNSUPPORTED_FILE


def test_empty_file_rejected():
    with pytest.raises(AppError) as ei:
        parse_document(b"", "a.pdf")
    assert ei.value.code == ErrorCode.BAD_PARAM


def test_docx_roundtrip():
    import docx

    buf = BytesIO()
    d = docx.Document()
    d.add_paragraph("姓名：张三")
    d.add_paragraph("硕士，5年工作经验")
    d.add_paragraph("技能：Python FastAPI MySQL")
    d.save(buf)
    result = parse_document(buf.getvalue(), "张三.docx")
    assert "张三" in result["text"]
    assert "FastAPI" in result["text"]


def test_pdf_roundtrip():
    from reportlab.pdfgen import canvas

    buf = BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(100, 700, "Hello Resume Text")
    c.drawString(100, 680, "Skills: Python")
    c.save()
    result = parse_document(buf.getvalue(), "resume.pdf")
    assert "Hello Resume Text" in result["text"]
    assert "Skills: Python" in result["text"]


def test_corrupt_pdf_rejected():
    with pytest.raises(AppError) as ei:
        parse_document(b"%PDF-1.4 garbage bytes", "broken.pdf")
    assert ei.value.code == ErrorCode.BAD_PARAM


def test_docx_textbox_content_extracted():
    """设计版简历内容在文本框（txbxContent）里，普通段落读不到——曾致此类简历 E4001 解析失败"""
    import docx
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    buf = BytesIO()
    d = docx.Document()
    d.add_paragraph("姓名：陈嘉航")
    # VML 文本框结构（与 Word 兼容存储的 Fallback 同构）；写两遍模拟 mc:Choice+Fallback 双份
    textbox_tpl = (
        '<w:p %s xmlns:v="urn:schemas-microsoft-com:vml">'
        "<w:r><w:pict><v:shape><v:textbox><w:txbxContent>"
        "<w:p><w:r><w:t>%%s</w:t></w:r></w:p>"
        "</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>"
    ) % nsdecls("w")
    for line in ("技能：Python FastAPI MySQL", "技能：Python FastAPI MySQL"):  # 同行×2 验证去重
        d.element.body.append(parse_xml(textbox_tpl % line))
    d.save(buf)

    result = parse_document(buf.getvalue(), "陈嘉航.docx")
    assert "陈嘉航" in result["text"]
    assert "文本框技能" not in result["text"]
    assert "技能：Python FastAPI MySQL" in result["text"]
    assert result["text"].count("技能：Python FastAPI MySQL") == 1  # 双份存储只计一次
