"""S01 文档解析：PDF/DOCX → 纯文本（对应 svc/contracts.py 的 parse_document）。

异常约定（契约）：
- AppError(E4002) 扩展名不支持 / 空文件名；
- AppError(E4001) 文件损坏或解析不出文本（如扫描件）。
"""
import re
from io import BytesIO
from pathlib import Path

from app.core.errors import AppError, ErrorCode

SUPPORTED_EXTS = {".pdf", ".docx"}
MAX_CHARS = 100_000  # 截断保护：超长文本会拖垮推理，简历场景足够


def parse_document(file_bytes: bytes, filename: str) -> dict:
    """入参文件二进制+原始文件名；出参 {"text": str}。异常契约见模块 docstring。"""
    ext = Path(filename or "").suffix.lower()
    if ext not in SUPPORTED_EXTS:
        raise AppError(ErrorCode.UNSUPPORTED_FILE, f"文件类型不支持: {ext or '未知'}（仅支持 PDF/DOCX）")
    if not file_bytes:
        raise AppError(ErrorCode.BAD_PARAM, "文件内容为空")

    text = _parse_pdf(file_bytes) if ext == ".pdf" else _parse_docx(file_bytes)
    text = text.strip()
    if not text:
        raise AppError(ErrorCode.BAD_PARAM, "未能从文件中解析出文本（可能是扫描件/图片型PDF）")
    return {"text": text[:MAX_CHARS]}


def _parse_pdf(raw: bytes) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(raw))
        pages = [(page.extract_text() or "") for page in reader.pages]
    except AppError:
        raise
    except Exception as exc:  # pypdf 各种损坏形态
        raise AppError(ErrorCode.BAD_PARAM, f"PDF 文件损坏或无法解析: {exc}") from exc
    return "\n".join(pages)


def _walk_table(table, parts: list[str]) -> None:
    """表格文本收集（C3 挑刺修复）：合并单元格只计一次、嵌套表格递归读出。

    python-docx 对横向合并（gridspan）的格会在 row.cells 里重复返回同一文本，
    原实现逐格追加 → 文本膨胀；嵌套表格（单元格里还有表）则整块丢失。
    """
    for row in table.rows:
        prev = None
        for cell in row.cells:
            text = cell.text
            if text.strip() and text != prev:  # 相邻重复 = 同一合并格，跳过
                parts.append(text)
            prev = text
            for nested in cell.tables:  # 递归：嵌套表此前整块漏读
                _walk_table(nested, parts)


def _parse_docx(raw: bytes) -> str:
    try:
        import docx
        from docx.oxml.ns import qn

        document = docx.Document(BytesIO(raw))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        # 简历常用表格排版，表格单元格文本同样计入
        for table in document.tables:
            _walk_table(table, parts)
        # 设计版简历的内容常整块放在文本框（txbxContent）里，普通段落读不到；
        # Word 兼容存储会把同一文本框写两遍（mc:Choice + mc:Fallback），按行去重
        seen: set[str] = set()
        for box in document.element.body.iter(qn("w:txbxContent")):
            for p in box.iter(qn("w:p")):
                line = "".join(t.text or "" for t in p.iter(qn("w:t"))).strip()
                if line and line not in seen:
                    seen.add(line)
                    parts.append(line)
    except AppError:
        raise
    except Exception as exc:
        raise AppError(ErrorCode.BAD_PARAM, f"DOCX 文件损坏或无法解析: {exc}") from exc
    return "\n".join(parts)


def looks_like_resume(text: str) -> bool:
    """轻量启发式：判断文本是否像一份简历（用于解析后质量提示，不阻断流程）。

    命中任意强特征即认为是简历：姓名字段 / 教育经历 / 工作经历 / 技能清单。
    """
    head = text[:2000]
    return bool(
        re.search(r"(姓名|教育|工作经历|项目经验|技能|求职意向|自我评价)", head)
    )
